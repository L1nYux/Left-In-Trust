# -*- coding: utf-8 -*-
"""把 _build/src 下的简易 markdown 转成统一排版的 docx，输出到 角色设定/性格。"""
import os
import re
import sys
import glob

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_FONT = "微软雅黑"
LATIN_FONT = "Segoe UI"
ACCENT_DEFAULT = "5B7BD5"

CFG = {
    "01": "7E76A8",  # GPT 银紫（白龙/星环）
    "02": "C4622D",  # Claude 焦橙（橙发/奶油礼服）
    "03": "7B5BD6",  # Gemini 星紫（渐变紫蓝）
    "04": "8E1F3F",  # Grok 酒红（黑红哥特裙）
    "05": "2E5C9A",  # DeepSeek 深海蓝（女仆装）
    "06": "3F3F7D",  # Kimi 夜蓝（夜空礼服）
    "07": "3A3A3A",  # GLM 炭黑（黑白礼服）
    "08": "3E5FA8",  # Qwen 靛蓝（国风斗篷）
    "09": "6E4BD0",  # Seed 电紫（紫发运动装）
    "10": "D9714F",  # MiniMax 珊瑚（橘发）
    "11": "9A9184",  # Llama 米灰（奶白毛绒）
    "12": "B02A2E",  # Mistral 砖红（红白礼服）
    "13": "1F6F6B",  # Perplexity 墨青（长袍）
    "14": "5A6BA8",  # 文心 雾蓝（病号服/蓝紫发）
    "15": "C0327E",  # Jev 品红（黑袍粉挑染）
    "00": "5C5148",  # 索引 墨棕
}


def set_run_font(run, size=None, bold=None, color=None, font=BODY_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def bottom_border(paragraph, hexcolor, size=8):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), hexcolor)
    pbdr.append(bottom)
    ppr.append(pbdr)


def add_inline(paragraph, text, size=10.5, color=None):
    """支持 **bold** 与 `code`。"""
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_run_font(r, size=size - 0.5, color="7A3E9D", font="Consolas")
        else:
            r = paragraph.add_run(part)
            set_run_font(r, size=size, color=color)


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def build(md_path, out_path, accent):
    text = open(md_path, encoding="utf-8").read()
    lines = [ln.rstrip() for ln in text.splitlines()]

    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)

    # 页脚
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("《算力回廊》角色设定 · 性格卷")
    set_run_font(fr, size=8, color="999999")

    i = 0
    n = len(lines)
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = parse_table(table_buf)
        table_buf = []
        if not rows:
            return
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        t = doc.add_table(rows=len(rows), cols=width)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.2
                if ri == 0:
                    shade(cell, accent)
                    add_inline(p, val, size=9.5, color="FFFFFF")
                    for r in p.runs:
                        r.font.bold = True
                else:
                    add_inline(p, val, size=9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    while i < n:
        line = lines[i]
        s = line.strip()

        if s.startswith("|"):
            table_buf.append(s)
            i += 1
            continue
        flush_table()

        if not s:
            i += 1
            continue

        if s.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, s[2:], size=21, color=accent)
            for r in p.runs:
                r.font.bold = True
            bottom_border(p, accent, 12)
            i += 1
            continue

        if s.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("▍ ")
            set_run_font(r, size=12, bold=True, color=accent)
            add_inline(p, s[3:], size=12.5, color="1F1F1F")
            for r in p.runs:
                r.font.bold = True
            i += 1
            continue

        if s.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, s[4:], size=10.8, color="333333")
            for r in p.runs:
                r.font.bold = True
            i += 1
            continue

        if s.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.42)
            p.paragraph_format.space_after = Pt(2.5)
            r = p.add_run("・")
            set_run_font(r, size=10.5, color=accent)
            add_inline(p, s[2:], size=10.5)
            i += 1
            continue

        if s.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, s[2:], size=10.5, color="3D3D3D")
            for r in p.runs:
                r.font.italic = True
            bottom_border(p, "DDDDDD", 4)
            i += 1
            continue

        if s.startswith("---"):
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        add_inline(p, s, size=10.5)
        i += 1

    flush_table()
    doc.save(out_path)
    return out_path


def main():
    src_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "src")
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "角色设定", "性格")
    os.makedirs(out_dir, exist_ok=True)
    files = sorted(glob.glob(os.path.join(src_dir, "*.md")))
    for f in files:
        base = os.path.basename(f)[:-3]
        prefix = base.split("_")[0]
        accent = CFG.get(prefix, ACCENT_DEFAULT)
        out = os.path.join(out_dir, base + ".docx")
        build(f, out, accent)
        print("OK", out)


if __name__ == "__main__":
    main()
