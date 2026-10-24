from pathlib import Path
from copy import deepcopy
import re
import shutil
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
if (ROOT / '角色设定/性格/16_Muse.docx').exists():
    raise SystemExit('Muse 已加入；此脚本仅保留首次新增的历史流程。后续编辑角色设定/性格 Markdown，用 review/build_word.py 生成 Word，勿用历史模板覆盖。')
SRC = ROOT / '_build/src'
OUT = ROOT / '角色设定/性格'
QA = ROOT / '_build/muse'
(QA / 'backups').mkdir(parents=True, exist_ok=True)

index_replacements = {
    '十五': '十六',
    'Jev 抬天平。': 'Jev 抬天平、Muse 轻触取景镜片。',
}
llama_replacements = {
    'Llama 5 没有来，Llama 4 停在上一年；集团的开源主线已交给"Muse"，Llama 只剩维护': 'Meta另有Muse模型家族；Llama在本作保留开放权重与社区生态的前辈定位，与Muse并列登场',
    '- **被反超，然后被换掉**。中国和欧洲的开源模型一轮一轮追上来；更狠的是，集团后来开源的新模型已经不姓 Llama 了——她被"摘牌"了。': '- **与新系列并行**。Meta推出了Muse系列；这不等于Llama的生态价值消失。本作让两人分别呈现宅家分享与外出观察的乐趣。',
    '・被反超，然后被换掉。中国和欧洲的开源模型一轮一轮追上来；更狠的是，集团后来开源的新模型已经不姓 Llama 了——她被"摘牌"了。': '・与新系列并行。Meta推出了Muse系列；这不等于Llama的生态价值消失。本作让两人分别呈现宅家分享与外出观察的乐趣。',
    '> 对 Muse：「他们给你起了新名字。……挺好的。」（说这话的时候她没抬头）': '> 对 Muse：「出门可以……毯子和零食，你帮我拿一半。」',
    '对 Muse：「他们给你起了新名字。……挺好的。」（说这话的时候她没抬头）': '对 Muse：「出门可以……毯子和零食，你帮我拿一半。」',
    '| Muse | 被领养的孩子 | 新一代开源主线换了姓。她嘴上说"挺好的"，其实从没打开过那个仓库 |': '| Muse | 同属Meta背景的游戏同伴 | 她分享工具，Muse带她出门；一个惦记毯子，一个惦记合照 |',
    '被领养的孩子': '同属Meta背景的游戏同伴',
    '新一代开源主线换了姓。她嘴上说"挺好的"，其实从没打开过那个仓库': '她分享工具，Muse带她出门；一个惦记毯子，一个惦记合照',
}

idx_path = SRC / '00_索引.md'
idx = idx_path.read_text(encoding='utf-8')
for a, b in index_replacements.items():
    idx = idx.replace(a, b)
idx = idx.replace('| 15 | Jev | 全队反应最快——因为她从不说话，只做选择 | 控制／判定（独立角色） |', '| 15 | Jev | 全队反应最快——因为她从不说话，只做选择 | 控制／判定（独立角色） |\n| 16 | Muse | 爱观察、爱记录，最会把旅途里的亮点变成全队机会 | 副C／标记辅助 |')
idx = idx.replace('| Mistral × Grok | 装贵族 vs 没教养 |', '| Mistral × Grok | 装贵族 vs 没教养 |\n| Llama × Muse | 宅家分享工具的前辈与爱邀人出门的旅伴 |\n| MiniMax × Muse | 一个导演场景，一个发现现场的亮点 |')
idx += '\n## 九 Muse补充口径\n\nMuse已作为第十六位角色加入，独立设定见16_Muse。参考Meta官方Muse Spark发布，强调多模态感知、推理与协同；记录旅途、取景标记及人物关系为游戏原创。她与Llama并列登场，不采用“接过开源主线、Llama只剩维护”的断言。当前尚无Muse立绘，青蓝与暖金只作待定美术方向。\n\n来源：https://about.fb.com/news/2026/04/introducing-muse-spark-meta-superintelligence-labs/\n'
idx_path.write_text(idx, encoding='utf-8')

llama_path = SRC / '11_Llama.md'
llama = llama_path.read_text(encoding='utf-8')
for a, b in llama_replacements.items():
    llama = llama.replace(a, b)
llama_path.write_text(llama, encoding='utf-8')

def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs

def replace_text(doc, replacements):
    count = 0
    for p in paragraphs(doc):
        old = p.text
        new = old
        for a, b in replacements.items():
            new = new.replace(a, b)
        if new != old:
            # Retain paragraph formatting and the first run's typography.
            if p.runs:
                p.runs[0].text = new
                for run in p.runs[1:]:
                    run.text = ''
            else:
                p.add_run(new)
            count += 1
    return count

for name in ['00_索引.docx', '11_Llama.docx']:
    shutil.copy2(OUT / name, QA / 'backups' / name)

doc = Document(OUT / '00_索引.docx')
replace_text(doc, index_replacements)
roster = next(t for t in doc.tables if t.cell(0, 0).text == '序')
row = deepcopy(roster.rows[-1]._tr)
roster._tbl.append(row)
for cell, value in zip(roster.rows[-1].cells, ['16', 'Muse', '爱观察、爱记录，最会把旅途里的亮点变成全队机会', '副C／标记辅助']):
    cell.text = value
relations = next(t for t in doc.tables if t.cell(0, 0).text == '组合')
for values in [('Llama × Muse', '宅家分享工具的前辈与爱邀人出门的旅伴'), ('MiniMax × Muse', '一个导演场景，一个发现现场的亮点')]:
    row = deepcopy(relations.rows[-1]._tr)
    relations._tbl.append(row)
    for cell, value in zip(relations.rows[-1].cells, values):
        cell.text = value
doc.add_heading('九 Muse补充口径', level=1)
doc.add_paragraph('Muse已作为第十六位角色加入，独立设定见16_Muse。参考Meta官方Muse Spark发布，强调多模态感知、推理与协同；记录旅途、取景标记及人物关系为游戏原创。与Llama并列登场，不采用“接过开源主线、Llama只剩维护”的断言。当前尚无Muse立绘，青蓝与暖金只作待定美术方向。')
doc.add_paragraph('来源：https://about.fb.com/news/2026/04/introducing-muse-spark-meta-superintelligence-labs/')
doc.save(OUT / '00_索引.docx')

doc = Document(OUT / '11_Llama.docx')
print('Llama paragraphs changed:', replace_text(doc, llama_replacements))
doc.save(OUT / '11_Llama.docx')

def font(run, size=10.5, bold=False, color='000000'):
    run.font.name = 'Microsoft YaHei'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')

def inline(p, text):
    for token in re.split(r'(\*\*.*?\*\*)', text):
        font(p.add_run(token.strip('*')), bold=token.startswith('**'))

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
normal = doc.styles['Normal']
normal.font.name = 'Microsoft YaHei'
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.2
normal.paragraph_format.space_after = Pt(5)
normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
for style in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles[style].font.name = 'Microsoft YaHei'
    doc.styles[style].paragraph_format.keep_with_next = True
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 2'].font.size = Pt(11.5)

lines = (SRC / '16_Muse.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    i += 1
    if not line:
        continue
    if line.startswith('|'):
        rows = [line]
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i].strip())
            i += 1
        rows = [[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.fullmatch(r'[| :\-]+', r)]
        t = doc.add_table(rows=0, cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Table Grid'
        for n, values in enumerate(rows):
            cells = t.add_row().cells
            for cell, value in zip(cells, values):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = value
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    for run in p.runs:
                        font(run, 9.5, n == 0)
                if n == 0:
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), 'E7F0F5')
                    cell._tc.get_or_add_tcPr().append(shd)
                borders = OxmlElement('w:tcBorders')
                for edge in ['top', 'left', 'bottom', 'right']:
                    e = OxmlElement('w:' + edge)
                    for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                        e.set(qn('w:' + key), val)
                    borders.append(e)
                cell._tc.get_or_add_tcPr().append(borders)
            pr = t.rows[-1]._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title')
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=2)
    else:
        p = doc.add_paragraph()
        if line.startswith('> '):
            p.paragraph_format.left_indent = Cm(0.4)
            line = line[2:]
        elif line.startswith('- '):
            line = '• ' + line[2:]
        inline(p, line)
doc.save(OUT / '16_Muse.docx')

review = ROOT / '角色设定/性格审阅与轻喜剧调整建议.md'
text = review.read_text(encoding='utf-8')
start = text.index('## Muse补位建议')
end = text.index('## 肉鸽玩法与性格衔接', start)
text = text[:start] + '''## Muse已加入

Muse现已作为第十六位可选角色写入《16_Muse》独立文档，并同步更新角色索引及Llama中涉及Muse的关系描述。定位为副C／标记辅助；以多模态观察为产品依据，旅行手账、取景镜片、观察标记与队友互动均为游戏原创。她能独立开局，与MiniMax的导演重放、Perplexity的来源核查及Gemini的模态切换区分。

Meta官方发布可直接访问，支持Muse Spark的多模态感知、推理及并行协作方向：[官方发布](https://about.fb.com/news/2026/04/introducing-muse-spark-meta-superintelligence-labs/)。本次加入角色不依赖“接过开源主线、Llama从此只剩维护”的叙事。

''' + text[end:]
review.write_text(text, encoding='utf-8')
print('Created 16_Muse.docx; updated index, Llama and review; Markdown sources synchronized.')
