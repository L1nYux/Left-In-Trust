from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
if (root / '角色设定/性格/00_索引.md').exists():
    raise SystemExit('历史排版脚本停用。当前 Markdown 请用 _build/review/build_word.py 生成。')
for name in ['16_Muse', '00_索引', '11_Llama']:
    path = root / '角色设定/性格' / (name + '.docx')
    doc = Document(path)
    # Remove title/quote border residue inherited from the Word template.
    for element in [doc._element, doc.styles.element]:
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for style in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    for n, p in enumerate(doc.paragraphs):
        if n == 0 or (p.runs and p.runs[0].text.startswith('▍')):
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        if '与新系列并行' in p.text:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
    for table in doc.tables:
        for row in list(table.rows):
            if all(re.fullmatch(r'[-: ]+', cell.text or ' ') for cell in row.cells):
                table._tbl.remove(row._tr)
        ncols = len(table.columns)
        total = 17 if name == '16_Muse' else 16.4
        if ncols == 2:
            widths = [3, total - 3]
        elif ncols == 3:
            widths = [2.5, 4, total - 6.5]
        elif ncols == 4 and table.cell(0, 0).text == '序':
            widths = [0.8, 2.1, 8.1, total - 11]
        else:
            widths = [2.8, 3.9, 3.9, total - 10.6]
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for i, row in enumerate(table.rows):
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = i == 0
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn('w:cantSplit')) is None:
                trpr.append(OxmlElement('w:cantSplit'))
            if i == 0 and trpr.find(qn('w:tblHeader')) is None:
                trpr.append(OxmlElement('w:tblHeader'))
        # A table header must stay with at least its first data row.
        if len(table.rows) > 1:
            for cell in table.rows[0].cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    if name == '16_Muse':
        doc.styles['Normal'].paragraph_format.line_spacing = 1.12
        doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    doc.save(path)
    print('Polished', name)

# Validate the added character and avoid accidental duplicate roster entries.
index = Document(root / '角色设定/性格/00_索引.docx')
roster = next(t for t in index.tables if t.cell(0, 0).text == '序')
assert len(roster.rows) == 17
assert roster.rows[-1].cells[1].text == 'Muse'
muse = Document(root / '角色设定/性格/16_Muse.docx')
assert len([p for p in muse.paragraphs if p.style.name == 'Heading 1']) == 9
assert 'Muse' in '\n'.join(p.text for p in muse.paragraphs)
print('Validated: 16 roster entries and 9 Muse sections')
